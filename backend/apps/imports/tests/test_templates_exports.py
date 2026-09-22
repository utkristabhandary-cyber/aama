"""TEMPLATE & EXPORT endpoints (Phase H).

Covers the admin-only template/export .xlsx endpoints:

- Structure: xlsx content-type, attachment name, sheet layout, header counts.
- Authorization: only admins may download templates/exports (backend authority).
- Unknown kinds are 404, mirroring the existing preview/confirm behavior.
- Honesty: templates are headers-only (uploading one yields the deterministic
  ``no_data_rows`` 400 — never a crash or fake rows).
- Round-trips: an export of real rows re-imports through the matching preview
  with no errors and a known plan (new / unchanged / update).
"""
import io
from datetime import date, time

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework import status

from apps.academics.models import Section, Semester, Subject, SubjectType, TimetableSlot
from apps.imports.contracts import (
    STUDENT_EXPORT_FILENAME,
    STUDENT_TEMPLATE_FILENAME,
    TEACHER_EXPORT_FILENAME,
    TEACHER_SYSTEM_HEADERS,
    TEACHER_TEMPLATE_FILENAME,
    TIMETABLE_EXPORT_FILENAME,
    TIMETABLE_SYSTEM_HEADERS,
    TIMETABLE_TEMPLATE_FILENAME,
    XLSX_CONTENT_TYPE,
)
from apps.imports.exports import (
    STUDENT_EXPORT_HEADERS,
    TEACHER_EXPORT_HEADERS,
    TIMETABLE_EXPORT_HEADERS,
)
from apps.imports.tests.helpers import ImportsBase
from apps.students.models import Student
from apps.teachers.models import Teacher

STUDENT_TEMPLATE_URL = "/api/imports/students/template/"
STUDENT_EXPORT_URL = "/api/imports/students/export/"
TEACHER_TEMPLATE_URL = "/api/imports/teachers/template/"
TEACHER_EXPORT_URL = "/api/imports/teachers/export/"
TIMETABLE_TEMPLATE_URL = "/api/academics/timetable-import/template/"
TIMETABLE_EXPORT_URL = "/api/academics/timetable-import/export/"


def workbook_of(response):
    return openpyxl.load_workbook(io.BytesIO(response.content))


class TemplateEndpointTests(ImportsBase):
    """Admin-only template downloads with correct structure and headers."""

    def test_student_template_structure(self):
        response = self.client.get(STUDENT_TEMPLATE_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        self.assertIn(f'filename="{STUDENT_TEMPLATE_FILENAME}"', response["Content-Disposition"])

        workbook = workbook_of(response)
        self.assertEqual(workbook.sheetnames, ["Students", "Import Contract"])
        header_row = list(workbook["Students"].iter_rows(values_only=True))
        self.assertEqual(len(header_row), 1)
        self.assertEqual(len(header_row[0]), 59)
        self.assertTrue(str(header_row[0][0]).startswith("S.N."))

        contract = list(workbook["Import Contract"].iter_rows(values_only=True))
        self.assertEqual(contract[0][0], "Column")
        self.assertGreater(len(contract), 60)

    def test_teacher_template_structure(self):
        response = self.client.get(TEACHER_TEMPLATE_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(f'filename="{TEACHER_TEMPLATE_FILENAME}"', response["Content-Disposition"])

        workbook = workbook_of(response)
        self.assertEqual(workbook.sheetnames, ["Teachers", "Import Contract"])
        headers = list(workbook["Teachers"].iter_rows(values_only=True))[0]
        self.assertEqual(headers, tuple(TEACHER_SYSTEM_HEADERS))

    def test_timetable_template_structure(self):
        response = self.client.get(TIMETABLE_TEMPLATE_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn(f'filename="{TIMETABLE_TEMPLATE_FILENAME}"', response["Content-Disposition"])

        workbook = workbook_of(response)
        self.assertEqual(workbook.sheetnames, ["Timetable", "Import Contract"])
        headers = list(workbook["Timetable"].iter_rows(values_only=True))[0]
        self.assertEqual(headers, tuple(TIMETABLE_SYSTEM_HEADERS))

    # -- authorization ------------------------------------------------

    def test_unauthenticated_template_rejected(self):
        self.clear_auth()
        for url in (STUDENT_TEMPLATE_URL, TEACHER_TEMPLATE_URL, TIMETABLE_TEMPLATE_URL):
            self.assertEqual(
                self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED, url
            )

    def test_teacher_role_rejected(self):
        self.auth_as(self.teacher)
        for url in (STUDENT_TEMPLATE_URL, TEACHER_TEMPLATE_URL, TIMETABLE_TEMPLATE_URL):
            self.assertEqual(
                self.client.get(url).status_code, status.HTTP_403_FORBIDDEN, url
            )

    def test_student_role_rejected(self):
        self.auth_as(self.student)
        for url in (STUDENT_TEMPLATE_URL, TEACHER_TEMPLATE_URL, TIMETABLE_TEMPLATE_URL):
            self.assertEqual(
                self.client.get(url).status_code, status.HTTP_403_FORBIDDEN, url
            )

    def test_unknown_kind_template_404(self):
        for url in ("/api/imports/classrooms/template/", "/api/imports/classrooms/export/"):
            response = self.client.get(url)
            self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND, url)
            self.assertEqual(response.data["code"], "kind_unknown")

    # -- blank template upload is a deterministic 400 -----------------

    def test_uploading_blank_student_template_yields_no_data_rows(self):
        template = self.client.get(STUDENT_TEMPLATE_URL).content
        response = self.client.post(
            "/api/imports/students/preview/",
            {"file": SimpleUploadedFile("students.xlsx", template)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data["code"], "no_data_rows")

    def test_uploading_blank_teacher_template_yields_no_data_rows(self):
        template = self.client.get(TEACHER_TEMPLATE_URL).content
        response = self.client.post(
            "/api/imports/teachers/preview/",
            {"file": SimpleUploadedFile("teachers.xlsx", template)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data["code"], "no_data_rows")

    def test_uploading_blank_timetable_template_yields_deterministic_400(self):
        # The timetable engine reports an empty workbook with its own code
        # ("INVALID_FILE" + "no timetable rows after the header row") rather
        # than the generic engine's "no_data_rows" — but the behaviour is the
        # same deterministic 400 with no crash and no fabricated rows.
        template = self.client.get(TIMETABLE_TEMPLATE_URL).content
        response = self.client.post(
            "/api/academics/timetable-import/preview/",
            {"file": SimpleUploadedFile("timetable.xlsx", template)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data["code"], "INVALID_FILE")
        self.assertIn("no timetable rows", response.data["detail"])


class ExportRoundTripTests(ImportsBase):
    """Exports reflect real rows and re-import with a clean, known plan."""

    def setUp(self):
        super().setUp()
        self.semester = Semester.objects.create(
            code="SEM-S4",
            name="Semester 4",
            academic_year="2025-2026",
            start_date=date(2026, 1, 12),
            end_date=date(2026, 5, 22),
            status="active",
        )
        self.section_a = Section.objects.create(
            semester=self.semester, name="A", capacity=60
        )
        self.section_b = Section.objects.create(
            semester=self.semester, name="B", capacity=60
        )
        self.subject = Subject.objects.create(
            semester=self.semester, code="CS101", name="Computer Science 101"
        )
        self.profile = Teacher.objects.create(
            teacher_id="tch-x1", name="Prof. Alpha", email="alpha@aams.local"
        )

    # -- student export -----------------------------------------------

    def test_student_export_headers_are_alias_friendly(self):
        Student.objects.create(
            student_id="STU-1",
            roll_no="R-1",
            name="Ada",
            email="ada@test.local",
            semester=self.semester,
            section=self.section_a,
        )
        response = self.client.get(STUDENT_EXPORT_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        headers = list(workbook_of(response)["Export"].iter_rows(values_only=True))[0]
        self.assertEqual(list(headers), STUDENT_EXPORT_HEADERS)
        # No column that would trip import validation or warnings.
        self.assertNotIn("Gender", headers)
        self.assertNotIn("DOB (B.S.)", headers)

    def test_student_export_round_trips(self):
        Student.objects.create(
            student_id="STU-1",
            roll_no="R-1",
            name="Ada",
            email="ada@test.local",
            phone="9800000000",
            dob=date(2004, 5, 12),
            admission_year=2025,
            semester=self.semester,
            section=self.section_a,
            address="Kathmandu",
            guardian_name="Ada Guardian",
            guardian_phone="9800000001",
            status="active",
        )
        exported = self.client.get(STUDENT_EXPORT_URL).content
        response = self.client.post(
            "/api/imports/students/preview/",
            {"file": SimpleUploadedFile("exported.xlsx", exported)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data["rows"]), 1)
        row = response.data["rows"][0]
        self.assertIn(row["plan"], {"new", "unchanged", "update"})
        self.assertFalse(
            any(i["level"] == "error" for i in row["issues"]),
            row["issues"],
        )

    def test_student_export_empty_db_emits_headers_only(self):
        response = self.client.get(STUDENT_EXPORT_URL)
        workbook = workbook_of(response)
        sheet = workbook["Export"]
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.max_column, len(STUDENT_EXPORT_HEADERS))

    # -- teacher export ------------------------------------------------

    def test_teacher_export_round_trips(self):
        response = self.client.get(TEACHER_EXPORT_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        workbook = workbook_of(response)
        headers = list(workbook["Export"].iter_rows(values_only=True))[0]
        self.assertEqual(list(headers), TEACHER_EXPORT_HEADERS)
        self.assertEqual(workbook["Export"].max_row, 2)

        exported = response.content
        response = self.client.post(
            "/api/imports/teachers/preview/",
            {"file": SimpleUploadedFile("exported.xlsx", exported)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        row = response.data["rows"][0]
        self.assertIn(row["plan"], {"new", "unchanged", "update"})
        self.assertFalse(any(i["level"] == "error" for i in row["issues"]))

    # -- timetable export ----------------------------------------------

    def _make_slots(self):
        combined = TimetableSlot.objects.create(
            semester=self.semester,
            section=self.section_a,
            subject=self.subject,
            teacher=self.profile,
            day="Monday",
            start_time=time(9, 0),
            end_time=time(10, 0),
            room="R-101",
            class_type=SubjectType.LECTURE,
        )
        combined.sections.set([self.section_a, self.section_b])
        TimetableSlot.objects.create(
            semester=self.semester,
            section=self.section_b,
            subject=self.subject,
            teacher=self.profile,
            day="Wednesday",
            start_time=time(11, 0),
            end_time=time(12, 0),
            room="R-102",
            class_type=SubjectType.TUTORIAL,
        )
        return combined

    def test_timetable_export_includes_combined_section_expression(self):
        self._make_slots()
        response = self.client.get(TIMETABLE_EXPORT_URL)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        workbook = workbook_of(response)
        headers = list(workbook["Export"].iter_rows(values_only=True))[0]
        self.assertEqual(list(headers), TIMETABLE_EXPORT_HEADERS)
        cells = [list(row) for row in workbook["Export"].iter_rows(values_only=True)][1:]
        sections = {c[1] for c in cells}
        self.assertIn("A+B", sections)

    def test_timetable_export_round_trips(self):
        slot = self._make_slots()
        exported = self.client.get(TIMETABLE_EXPORT_URL).content
        response = self.client.post(
            "/api/academics/timetable-import/preview/",
            {"file": SimpleUploadedFile("exported.xlsx", exported)},
            format="multipart",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        plans = [row["plan"] for row in response.data["rows"]]
        self.assertEqual(len(plans), 2)
        for row in response.data["rows"]:
            self.assertIn(row["plan"], {"new", "unchanged", "update"})
            self.assertFalse(any(i["level"] == "error" for i in row["issues"]))

    # -- export authorization -------------------------------------------

    def test_export_endpoints_admin_only(self):
        self.clear_auth()
        for url in (STUDENT_EXPORT_URL, TEACHER_EXPORT_URL, TIMETABLE_EXPORT_URL):
            self.assertEqual(
                self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED, url
            )
        self.auth_as(self.teacher)
        for url in (STUDENT_EXPORT_URL, TEACHER_EXPORT_URL, TIMETABLE_EXPORT_URL):
            self.assertEqual(
                self.client.get(url).status_code, status.HTTP_403_FORBIDDEN, url
            )

    def test_export_content_disposition_sets_filename(self):
        self.client.get(TEACHER_EXPORT_URL)
        response = self.client.get(TEACHER_EXPORT_URL)
        self.assertIn(f'filename="{TEACHER_EXPORT_FILENAME}"', response["Content-Disposition"])
        self.assertIn(
            f'filename="{TIMETABLE_EXPORT_FILENAME}"',
            self.client.get(TIMETABLE_EXPORT_URL)["Content-Disposition"],
        )